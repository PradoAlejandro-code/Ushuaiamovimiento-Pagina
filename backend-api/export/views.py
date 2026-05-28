import csv
import zipfile
import io
import os
from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from surveys.models import Encuesta, RespuestaHeader
from contact.models import Contacto
from django.db.models import Q

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter


from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def aplicar_estilo_header_excel(ws, titulo_principal, headers, base_dir):
    orange_fill = PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid")
    blue_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    white_bold_font = Font(bold=True, color="FFFFFF", size=12)
    header_font = Font(bold=True, color="FFFFFF")
    
    left_center_alignment = Alignment(horizontal="left", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions['A'].width = 50 

    # FILA 1: LOGO
    ws.row_dimensions[1].height = 58 
    ws.append([""]) 
    
    path_logo = os.path.join(base_dir, 'static', 'images', 'logo_excel.png')
    if os.path.exists(path_logo):
        img_logo = OpenpyxlImage(path_logo)
        img_logo.height = 75 
        img_logo.width = 225
        ws.add_image(img_logo, 'A1')

    # FILA 2: TITULO PRINCIPAL
    ws.row_dimensions[2].height = 25
    ws.append([f"  {titulo_principal}"]) 
    cell_nombre = ws.cell(row=2, column=1)
    cell_nombre.fill = orange_fill
    cell_nombre.font = white_bold_font
    cell_nombre.alignment = left_center_alignment

    # FILA 3: HEADERS
    ws.append(headers)
    for cell in ws[3]:
        cell.font = header_font
        cell.fill = blue_fill
        cell.alignment = center_alignment

class ExportarInformeDocxView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        informe = get_object_or_404(Informe, pk=pk)
        
        doc = Document()

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(14)

        # --- TÍTULO ---
        titulo = doc.add_heading(informe.titulo.upper(), 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # --- DESCRIPCIÓN ---
        doc.add_heading("Descripción", level=1)
        if informe.descripcion_breve:
            descripcion = doc.add_paragraph(informe.descripcion_breve)
            descripcion.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_paragraph(" ").alignment = WD_ALIGN_PARAGRAPH.CENTER

        # --- CUERPO ---
        doc.add_heading("Informe", level=1)
        cuerpo = doc.add_paragraph(informe.cuerpo)
        cuerpo.alignment = WD_ALIGN_PARAGRAPH.BOTH

        # --- ANEXO DE FOTOGRAFÍAS ---
        fotos = informe.fotos.all()
        if fotos.exists():
            doc.add_page_break()
            doc.add_heading("Anexo de Fotografías", level=1)
            
            for foto in fotos:
                try:
                    if foto.imagen and os.path.exists(foto.imagen.path):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r = p.add_run()
                        r.add_picture(foto.imagen.path, width=Inches(5.0))
                        caption = doc.add_paragraph(f"Referencia: {informe.titulo}")
                        caption.style = 'Caption'
                        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    print(f"Error al insertar foto {foto.id}: {e}")
                    continue

        # --- GENERAR RESPUESTA ---
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = f"Informe_{informe.id}.docx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
        
class ExportarContactosView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        # --- BUSQUEDA POR FILTROS ---
        search_query = request.query_params.get('search', '')
        tag_query = request.query_params.get('tag', '')

        contactos = Contacto.objects.all().order_by('nombre')

        if search_query:
            contactos = contactos.filter(
                Q(nombre__icontains=search_query) |
                Q(celular__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(dni__icontains=search_query) |
                Q(tag__icontains=search_query)
            )

        if tag_query:
            contactos = contactos.filter(tag__icontains=tag_query)

        # --- CSV CONFIGURACION ---
        output = io.StringIO()
        writer = csv.writer(output)
        
        headers = [
            "Nombre", 
            "Telefono", 
            "Mail", 
            "dni",
            "tag"
        ]
        writer.writerow(headers)

        for c in contactos:
            row = [
                c.nombre or "Sin nombre",
                c.celular or "",
                c.email or "",
                c.dni or "",
                c.tag or ""
            ]
            writer.writerow(row)

        # --- GENERAR RESPUESTA ---
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="contactos_import_google.csv"'
        
        return response

class ExportarContactosExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        search_query = request.query_params.get('search', '')
        tag_query = request.query_params.get('tag', '')

        contactos = Contacto.objects.all().order_by('nombre')

        if search_query:
            contactos = contactos.filter(
                Q(nombre__icontains=search_query) |
                Q(celular__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(dni__icontains=search_query) |
                Q(tag__icontains=search_query)
            )

        if tag_query:
            contactos = contactos.filter(tag__icontains=tag_query)

        wb = Workbook()
        ws = wb.active
        ws.title = "Contactos"

        max_tags = 1
        for c in contactos:
            if c.tag:
                num_tags = len([t for t in c.tag.split(',') if t.strip()])
                if num_tags > max_tags:
                    max_tags = num_tags

        headers = ["Nombre", "Telefono", "Mail", "DNI"]
        for i in range(max_tags):
            headers.append(f"Etiqueta {i+1}")
        
        aplicar_estilo_header_excel(ws, "CONTACTOS: Exportación General", headers, settings.BASE_DIR)

        data_alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

        for c in contactos:
            row_data = [
                c.nombre or "Sin nombre",
                c.celular or "",
                c.email or "",
                c.dni or ""
            ]
            
            tags_list = [t.strip() for t in c.tag.split(',') if t.strip()] if c.tag else []
            for i in range(max_tags):
                if i < len(tags_list):
                    row_data.append(tags_list[i])
                else:
                    row_data.append("")

            ws.append(row_data)
            for cell in ws[ws.max_row]:
                cell.alignment = data_alignment

        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            if column_letter == 'A': continue
            for cell in col:
                if cell.row < 4: continue
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        response = HttpResponse(
            excel_buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="Contactos.xlsx"'
        return response

class ExportarEncuestaCompletaView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        encuesta = get_object_or_404(Encuesta, pk=pk)
        
        # --- CAPTURA DE FILTROS ---
        secciones_filtradas = request.query_params.getlist('seccion') or request.query_params.getlist('seccion[]')
        barrios_filtrados = request.query_params.getlist('barrio') or request.query_params.getlist('barrio[]')
        
        con_fotos_param = request.query_params.get('con_fotos', 'true').lower()
        con_fotos = con_fotos_param in ['true', '1', 'yes']
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Respuestas"

        data_alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

        preguntas = encuesta.preguntas.filter(activa=True).order_by('orden')
        headers_static = ['ID Respuesta', 'Fecha', 'Usuario', 'Barrio', 'Seccion']
        
        headers_dinamicos = [p.titulo for p in preguntas]
        full_headers = headers_static + headers_dinamicos

        # Aplicar diseño global con helper
        aplicar_estilo_header_excel(ws, f"ENCUESTA: {encuesta.nombre}", full_headers, settings.BASE_DIR)

        # --- PREPARAR ZIP ---
        zip_buffer = io.BytesIO() if con_fotos else None
        zip_file = zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) if con_fotos else None

        # --- PROCESAMIENTO DE DATOS ---
        respuestas = RespuestaHeader.objects.filter(encuesta=encuesta).select_related('usuario').prefetch_related('detalles__pregunta', 'detalles__fotos_extra')

        if secciones_filtradas:
            respuestas = respuestas.filter(seccion__in=secciones_filtradas)

        if barrios_filtrados:
            respuestas = respuestas.filter(barrio__in=barrios_filtrados)

        for rta in respuestas:
            row_data = [
                rta.id,
                rta.fecha_envio.strftime('%Y-%m-%d %H:%M'),
                rta.usuario.username if rta.usuario else 'Anónimo',
                rta.barrio or '-',
                rta.seccion or '-'
            ]

            detalles_map = {d.pregunta_id: d for d in rta.detalles.all()}
            for p in preguntas:
                detalle = detalles_map.get(p.id)
                val = ""
                if detalle:
                    val = detalle.valor_texto or (str(detalle.valor_numero) if detalle.valor_numero is not None else "")
                    if p.tipo == 'foto':
                        fotos = detalle.fotos_extra.all()
                        if fotos.exists():
                            nombres_fotos = []
                            for idx, f in enumerate(fotos):
                                try:
                                    if f.imagen and os.path.exists(f.imagen.path):
                                        ext = f.imagen.name.split('.')[-1]
                                        zip_filename = f"imagenes/R{rta.id}_P{p.orden}_{idx+1}.{ext}"
                                        
                                        if con_fotos and zip_file:
                                            with open(f.imagen.path, 'rb') as img_f:
                                                zip_file.writestr(zip_filename, img_f.read())
                                                
                                        nombres_fotos.append(zip_filename)
                                except Exception: continue
                            if nombres_fotos: val = " | ".join(nombres_fotos)
                row_data.append(val)

            ws.append(row_data)
            
            for cell in ws[ws.max_row]:
                cell.alignment = data_alignment

        # --- AUTO-AJUSTE DE COLUMNAS ---
        for i, column_cells in enumerate(ws.columns, start=1):
            column_letter = get_column_letter(i)
            if column_letter == 'A': continue 
            
            max_length = 0
            for cell in column_cells:
                if cell.row < 4: continue 
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        # --- GENERAR RESPUESTAS FINAL ---
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        
        if con_fotos and zip_file and zip_buffer:
            zip_file.writestr('respuestas.xlsx', excel_buffer.getvalue())
            zip_file.close()

            zip_buffer.seek(0)
            response = HttpResponse(zip_buffer, content_type='application/zip')
            response['Content-Disposition'] = f'attachment; filename="Export_Encuesta_{encuesta.id}.zip"'
            return response
        else:
            excel_buffer.seek(0)
            response = HttpResponse(
                excel_buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="Export_Encuesta_{encuesta.id}.xlsx"'
            return response